"""Update the Note / check columns of the all-timestep rows after official_all.py ran."""

import json
from pathlib import Path

import openpyxl

from finalize import OUT

HERE = Path(__file__).parent

for f in sorted((HERE / "results_all").glob("*.json")):
    r = json.load(open(f))
    if "official_ok_all_timesteps" not in r:
        continue
    lk, v = r["leveltype"], r["variable"]
    path = OUT / ("ERA5-Pressure.xlsx" if lk == "pressure" else "ERA5-Single.xlsx")
    wb = openpyxl.load_workbook(path)
    ws, ws2 = wb.worksheets[0], wb["vs-scoreboard"]
    hdr = [c.value for c in ws[1]]
    label = f"all ({r['T']} timesteps, each compressed independently)"
    status = "official check_safety_requirements passed on ALL timesteps" if r["official_ok_all_timesteps"] else f"official checker FAILED on timesteps {r['official_failed_timesteps']}"
    note = f"ratio = total original bytes / total compressed bytes over {r['T']} timesteps; per-timestep ratio {r['cr_step_min']:.2f}-{r['cr_step_max']:.2f}; {status}"
    for row in ws.iter_rows(min_row=2):
        if row[hdr.index("Variable")].value == v and row[hdr.index("Default/all timestep(s)")].value == label:
            row[hdr.index("Note")].value = note
    h2 = [c.value for c in ws2[1]]
    for row in ws2.iter_rows(min_row=2):
        if row[0].value == f"{v} [{label}]":
            row[h2.index("Official check")].value = "ok (all timesteps)" if r["official_ok_all_timesteps"] else "FAILED"
            row[h2.index("Note")].value = note
    wb.save(path)
    print(lk, v, status)
