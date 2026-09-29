"""Append all-timestep results (results_all/*.json) as extra rows to the ERA5 Excel files."""

import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment

from finalize import AUTHOR, OUT, REC_VERSION, explicit_code

HERE = Path(__file__).parent

for f in sorted((HERE / "results_all").glob("*.json")):
    r = json.load(open(f))
    if "cr_all" not in r:
        continue
    lk, v = r["leveltype"], r["variable"]
    path = OUT / ("ERA5-Pressure.xlsx" if lk == "pressure" else "ERA5-Single.xlsx")
    wb = openpyxl.load_workbook(path)
    ws, ws2 = wb.worksheets[0], wb["vs-scoreboard"]
    hdr = [c.value for c in ws[1]]
    label = f"all ({r['T']} timesteps, each compressed independently)"
    # skip if already present
    if any(row[hdr.index("Variable")] == v and row[hdr.index("Default/all timestep(s)")] == label for row in ws.iter_rows(min_row=2, values_only=True)):
        continue
    code = explicit_code(r["config"])
    short = r["config_short"] + ("" if not r["changed"] else " (parameter tightened for all timesteps)")
    note = f"ratio = total original bytes / total compressed bytes over {r['T']} timesteps; per-timestep ratio {r['cr_step_min']:.2f}-{r['cr_step_max']:.2f}; requirement checked on every timestep (fast mirror checker) and with the official checker on timesteps {list(r['official_ok_samples'])}"
    enc, dec = round(r["enc_gbps"], 4), round(r["dec_gbps"], 4)
    if lk == "pressure":
        row = [AUTHOR, round(r["cr_all"], 2), dec, v, label, short, REC_VERSION, repr(r["config"]), enc, dec, code, note]
    else:
        row = [AUTHOR, round(r["cr_all"], 2), v, label, short, REC_VERSION, repr(r["config"]), enc, dec, code, note]
    ws.append(row)
    for c in ws[ws.max_row]:
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws2.append([f"{v} [{label}]", r["requirements"], round(r["cr_all"], 2), None, None, None, f"default-timestep ratio of the same config: {r['cr_default_timestep']:.2f}",
                "ok" if all(r["official_ok_samples"].values()) else "FAILED", None, None, note])
    wb.save(path)
    print(lk, v, "CR_all", round(r["cr_all"], 2), "changed", r["changed"])
