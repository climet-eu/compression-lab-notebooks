"""
Re-verify all winning configurations (official checks, exact notebook checks
for the three small challenges), measure throughput, compare with the current
scoreboard, and write one Excel file per challenge.

Usage:  python finalize.py [nproc]
"""

import json
import sys
import time
import traceback
from pathlib import Path

import numpy as np
import openpyxl
from numcodecs.registry import get_codec
from openpyxl.styles import Font

import numcodecs_interp_ctx  # noqa: F401
import numcodecs_lon_gradient  # noqa: F401
import numcodecs_chunked  # noqa: F401
import numcodecs_clip  # noqa: F401
import numcodecs_context_mixing  # noqa: F401
import numcodecs_eb_quantize  # noqa: F401
import numcodecs_abs_or_rel  # noqa: F401
import numcodecs_grid_int  # noqa: F401
import numcodecs_mask  # noqa: F401
import numcodecs_pw_ratio  # noqa: F401
import numcodecs_replace  # noqa: F401
import numcodecs_zero  # noqa: F401
from numcodecs_context_mixing import ContextMixingBitmapCodec, ContextMixingResidualCodec, ContextMixingSymbolCodec
from numcodecs_eb_quantize import ErrorBoundedQuantizeCodec
from numcodecs_mask import MaskMetaCodec
from numcodecs_pw_ratio import PointwiseRatioErrorBoundedCodec
from numcodecs_lon_gradient import LongitudeGradientCodec

HERE = Path(__file__).parent
OUT = HERE / "submissions"
OUT.mkdir(exist_ok=True)
AUTHOR = "@SF-N"
REC_VERSION = "0.1.0a2"
NOTE = "codecs: SF-N/numcodecs-* packages (clip, chunked, grid-int, eb-quantize, context-mixing, interp-ctx, lon-gradient, abs-or-rel), juntyr/numcodecs-mask (PR #4), numcodecs-replace (ThresholdFilterCodec PR), numcodecs-zero (value PR), numcodecs-pw-ratio"

SMALL_HEADER = [
    "Author",
    "Compression",
    "Config Short",
    "Configuration",
    "[opt] Throughput Compression [GB/s], measured relative to the original size",
    "[opt] Throughput Decompression [GB/s], measured relative to the original size",
]
PRESSURE_HEADER = [
    "Author",
    "Compression",
    "Decompression Throughput (GB/s)",
    "Variable",
    "Default/all timestep(s)",
    "Config Short",
    "Version of compression-recommendations",
    "Configuration",
    "[opt] Throughput Compression [GB/s], measured relative to the original size",
    "[opt] Throughput Decompression [GB/s], measured relative to the original size",
]
SINGLE_HEADER = [
    "Author",
    "Compression",
    "Variable",
    "Default/all timestep(s)",
    "Config Short",
    "Version of compression-recommendations",
    "Configuration",
    "[opt] Throughput Compression [GB/s], measured relative to the original size",
    "[opt] Throughput Decompression [GB/s], measured relative to the original size",
]


def timed_roundtrip(codec, x):
    t = time.perf_counter()
    e = codec.encode(x)
    te = time.perf_counter() - t
    buf = np.frombuffer(bytes(e), np.uint8)
    t = time.perf_counter()
    d = codec.decode(buf)
    td = time.perf_counter() - t
    d = np.asarray(d).reshape(x.shape)
    return e, d, te, td


def config_str(codec):
    return repr(codec.get_config())


# id -> (module, class); nested "inner"/"lossless" configs become nested constructor calls
CLASSES = {
    "interp_ctx": ("numcodecs_interp_ctx", "InterpolationContextMixingCodec"),
    "lon_gradient": ("numcodecs_lon_gradient", "LongitudeGradientCodec"),
    "clip": ("numcodecs_clip", "ClipCodec"),
    "abs_or_rel": ("numcodecs_abs_or_rel", "AbsOrRelErrorBoundedCodec"),
    "grid_int": ("numcodecs_grid_int", "GridIntCodec"),
    "chunked": ("numcodecs_chunked", "ChunkedCodec"),
    "combinators.stack": ("numcodecs_combinators.stack", "CodecStack"),
    "eb_quantize": ("numcodecs_eb_quantize", "ErrorBoundedQuantizeCodec"),
    "context_mixing.bitmap": ("numcodecs_context_mixing", "ContextMixingBitmapCodec"),
    "context_mixing.symbols": ("numcodecs_context_mixing", "ContextMixingSymbolCodec"),
    "context_mixing.residuals": ("numcodecs_context_mixing", "ContextMixingResidualCodec"),
    "mask.meta": ("numcodecs_mask", "MaskMetaCodec"),
    "replace.filter": ("numcodecs_replace", "ReplaceFilterCodec"),
    "replace.threshold": ("numcodecs_replace", "ThresholdFilterCodec"),
    "zero": ("numcodecs_zero", "ZeroCodec"),
    "sperr.rs": ("numcodecs_wasm_sperr", "Sperr"),
    "sz3.rs": ("numcodecs_wasm_sz3", "Sz3"),
    "zfp.rs": ("numcodecs_wasm_zfp", "Zfp"),
    "zstd.rs": ("numcodecs_wasm_zstd", "Zstd"),
    "lzma": ("numcodecs", "LZMA"),
    "pw_ratio": ("numcodecs_pw_ratio", "PointwiseRatioErrorBoundedCodec"),
}


def explicit_code(config):
    """Fully explicit Python: imports + nested constructor calls for `config`."""
    imports = set()

    def literal(v):
        if isinstance(v, float) and v != v:
            return 'float("nan")'
        if isinstance(v, float) and v in (float("inf"), float("-inf")):
            return 'float("inf")' if v > 0 else 'float("-inf")'
        return repr(v)

    def expr(cfg, indent):
        cid = cfg["id"]
        mod, cls = CLASSES[cid]
        imports.add(f"from {mod} import {cls}")
        pad = " " * indent
        args = []
        if cid == "combinators.stack":
            for v in cfg["codecs"]:
                args.append(f"{pad}    {expr(v, indent + 4)},")
            return f"{cls}(\n" + "\n".join(args) + f"\n{pad})"
        for k, v in cfg.items():
            if k in ("id", "_version"):
                continue
            if isinstance(v, dict) and "id" in v and not (cid in ("pw_ratio", "lon_gradient", "abs_or_rel") and k in ("log_codec", "codec")):
                args.append(f"{pad}    {k}={expr(v, indent + 4)},")
            else:
                args.append(f"{pad}    {k}={literal(v)},")
        if not args:
            return f"{cls}()"
        return f"{cls}(\n" + "\n".join(args) + f"\n{pad})"

    body = expr(config, 0)
    return "\n".join(sorted(imports)) + "\n\ncodec = " + body




def codec_code(codec, direct=None):
    """Python snippet reproducing `codec` in a challenge notebook."""
    return explicit_code(codec.get_config())


# ----------------------------------------------------------------------------
# small challenges (exact notebook checks)
# ----------------------------------------------------------------------------
def run_nan():
    import xarray as xr

    ds = xr.open_dataset(HERE / "data/HOAPS_2020-08_6-hourly.nc", engine="h5netcdf", decode_timedelta=True)
    da = ds["wvpa"].sel(time=slice("2020-08-01", "2020-08-07"))
    codec = MaskMetaCodec(
        mask=float("nan"),
        bitmap_codec=ContextMixingBitmapCodec(mixer_rate=0.005),
        codec=ErrorBoundedQuantizeCodec(codec=ContextMixingSymbolCodec(), eb=1.0),
    )
    direct = None
    e, d, te, td = timed_roundtrip(codec, da.values)
    da_dec = da.copy(data=d)
    violations = float(np.mean(xr.where(np.isnan(da), ~np.isnan(da_dec), ~(np.abs(da_dec - da) <= 1))))
    cr = da.nbytes / np.array(e).nbytes
    return dict(cr=cr, violations=violations, te=te, td=td, nbytes=da.nbytes, codec=codec, direct=direct,
                short="mask.meta(NaN, bitmap=context_mixing.bitmap) + eb_quantize(eb=1) + context_mixing.symbols")


def run_pwrel():
    import xarray as xr

    ds = xr.open_dataset(HERE / "data/hplp_sfc_regridded_tp_025deg_steps_228_240.nc", engine="h5netcdf", decode_timedelta=True)
    da = ds["tp"]
    codec = PointwiseRatioErrorBoundedCodec(
        eb_ratio=1.01,
        eb_abs_marker="$eb_abs",
        log_codec={"id": "eb_quantize", "eb": "$eb_abs", "codec": {"id": "context_mixing.residuals", "mixer_rate": 0.003}},
        sign_codec={"id": "zstd.rs", "level": 19},
    )
    direct = None
    e, d, te, td = timed_roundtrip(codec, da.values)
    da_dec = da.copy(data=d)
    violations = float(np.mean(~(np.abs(da_dec - da) <= (np.abs(da) * 0.01))))
    cr = da.nbytes / np.array(e).nbytes
    return dict(cr=cr, violations=violations, te=te, td=td, nbytes=da.nbytes, codec=codec, direct=direct,
                short="pw_ratio(1.01) + eb_quantize + context_mixing.residuals")


def run_gradient():
    import xarray as xr

    ds = xr.open_dataset(HERE / "data/NextGEMS_regridded_hus_025deg_steps_44_45.nc", engine="h5netcdf", decode_timedelta=True)
    da = ds["hus"]
    codec = LongitudeGradientCodec(
        eb=1e-6, spacing=0.25, stencil=5, shrink=0.02,
        codec={"id": "eb_quantize", "eb": "$eb_abs", "codec": {"id": "context_mixing.residuals"}},
    )
    direct = None
    e, d, te, td = timed_roundtrip(codec, da.values)
    da_dec = da.copy(data=d)

    def deriv(a):
        return (a.roll(lon=-5) - a.roll(lon=5)) / (np.mod(a.lon.roll(lon=-5) - a.lon.roll(lon=5), 360))

    violations = float(np.mean(~(np.abs(deriv(da_dec) - deriv(da)) <= 1e-6)))
    cr = da.nbytes / np.array(e).nbytes
    return dict(cr=cr, violations=violations, te=te, td=td, nbytes=da.nbytes, codec=codec, direct=direct,
                short="lon_gradient: stride-10 longitude differences (eb 2.5e-6, closure-aware rounding) + eb_quantize + context_mixing.residuals")


def write_small(name, r, prev_best):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = name
    ws.append(SMALL_HEADER + ["Codec (paste into the notebook's compressor cell)"])
    ws.append([AUTHOR, round(r["cr"], 2), r["short"], config_str(r["codec"]), round(r["nbytes"] / r["te"] / 1e9, 4), round(r["nbytes"] / r["td"] / 1e9, 4), codec_code(r["codec"], r["direct"])])
    ws2 = wb.create_sheet("vs-scoreboard")
    ws2.append(["Challenge", "New compression ratio", "Previous best", "Improvement", "Violations", "Note"])
    ws2.append([name, round(r["cr"], 2), prev_best, f"{(r['cr'] / prev_best - 1) * 100:+.1f}%", r["violations"], NOTE])
    for w in (ws, ws2):
        for c in w[1]:
            c.font = Font(bold=True)
    wb.save(OUT / f"{name}.xlsx")


# ----------------------------------------------------------------------------
# ERA5
# ----------------------------------------------------------------------------
def verify_era5(args):
    lk, v = args
    try:
        res = json.load(open(HERE / "results" / f"{lk}__{v}.json"))
        if "config" not in res:
            return (lk, v, None)
        x = np.load(HERE / "data/era5" / f"{lk}__{v}.npz")["arr"]
        from compression_requirement_checks import check_safety_requirements
        from search import get_requirements

        reqs = get_requirements(lk, v)
        codec = get_codec(res["config"])
        e, d, te, td = timed_roundtrip(codec, x)
        ok = bool(check_safety_requirements(original=x, reconstructed=d, requirements=reqs))
        cr = x.nbytes / len(e)
        fin = np.isfinite(x)
        err = (d[fin] - x[fin]).astype(np.float64)
        std = float(x[fin].std()) if fin.any() else 0.0
        nrmse = float(np.sqrt(np.mean(err ** 2)) / std) if std > 0 else 0.0
        maxerr = float(np.abs(err).max()) if err.size else 0.0
        return (lk, v, dict(cr=cr, ok=ok, te=te, td=td, nbytes=x.nbytes, size=len(e), config=config_str(codec), code=codec_code(codec),
                            short=res["config_short"], requirements=res["requirements"], nrmse=nrmse, maxerr=maxerr))
    except Exception:
        traceback.print_exc()
        return (lk, v, {"error": traceback.format_exc()})


def write_era5(lk, rows, best):
    wb = openpyxl.Workbook()
    ws = wb.active
    header = PRESSURE_HEADER if lk == "pressure" else SINGLE_HEADER
    ws.title = "ERA5-Pressure" if lk == "pressure" else "ERA5-Single"
    ws.append(header + ["Codec (paste into the notebook's compressor cell)", "Note"])
    ws2 = wb.create_sheet("vs-scoreboard")
    ws2.append(["Variable", "Requirements", "New compression ratio", "Previous best", "Previous author", "Previous config", "Improvement", "Official check", "NRMSE (rmse/std)", "Max abs error", "Note"])
    for v, r in sorted(rows.items()):
        if r is None or "error" in r:
            continue
        comp_t = round(r["nbytes"] / r["te"] / 1e9, 4)
        dec_t = round(r["nbytes"] / r["td"] / 1e9, 4)
        pb = best.get(f"{lk}__{v}")
        prev = pb[0] if pb else None
        note = NOTE
        flag = ""
        if r["nrmse"] > 0.5:
            flag = "DEGENERATE: the recommended requirement is very loose, reconstruction loses most of the field structure (NRMSE %.2f)." % r["nrmse"]
            note = flag + " " + NOTE
        if lk == "pressure":
            ws.append([AUTHOR, round(r["cr"], 2), dec_t, v, "default", r["short"], REC_VERSION, r["config"], comp_t, dec_t, r["code"], flag])
        else:
            ws.append([AUTHOR, round(r["cr"], 2), v, "default", r["short"], REC_VERSION, r["config"], comp_t, dec_t, r["code"], flag])
        ws2.append([v, r["requirements"], round(r["cr"], 2), prev, pb[1] if pb else None, pb[2] if pb else None,
                    (f"{(r['cr'] / prev - 1) * 100:+.1f}%" if prev else "new"), "ok" if r["ok"] else "FAILED",
                    round(r["nrmse"], 4), r["maxerr"], note])
    for w in (ws, ws2):
        for c in w[1]:
            c.font = Font(bold=True)
    wb.save(OUT / ("ERA5-Pressure.xlsx" if lk == "pressure" else "ERA5-Single.xlsx"))


if __name__ == "__main__":
    import multiprocessing as mp

    nproc = int(sys.argv[1]) if len(sys.argv) > 1 else 4
    which = sys.argv[2].split(",") if len(sys.argv) > 2 else ["small", "pressure", "single"]
    if not (HERE / "scoreboard_best.json").exists():
        import fetch_scoreboard

        fetch_scoreboard.main()
    best = json.load(open(HERE / "scoreboard_best.json"))
    summary = {}
    if "small" in which:
        for name, fn, prev in [("NaN", run_nan, 68.96), ("PwRel", run_pwrel, 21.76), ("Gradient", run_gradient, 9.11)]:
            r = fn()
            print(name, "CR", round(r["cr"], 2), "violations", r["violations"], "prev", prev, flush=True)
            write_small(name, r, prev)
            summary[name] = dict(cr=r["cr"], prev=prev, violations=r["violations"])
    for lk in ("pressure", "single"):
        if lk not in which:
            continue
        files = sorted((HERE / "results").glob(f"{lk}__*.json"))
        tasks = [(lk, f.stem.split("__", 1)[1]) for f in files]
        rows = {}
        with mp.get_context("spawn").Pool(nproc, maxtasksperchild=4) as pool:
            for lk_, v, r in pool.imap_unordered(verify_era5, tasks):
                rows[v] = r
                if r and "error" not in r:
                    pb = best.get(f"{lk}__{v}")
                    print(lk, v, round(r["cr"], 2), "prev", pb[0] if pb else None, "ok" if r["ok"] else "FAILED", flush=True)
                else:
                    print(lk, v, "ERROR/none", flush=True)
        write_era5(lk, rows, best)
        summary[lk] = {v: (r["cr"] if r and "error" not in r else None) for v, r in rows.items()}
    json.dump(summary, open(OUT / "summary.json", "w"), indent=1)
